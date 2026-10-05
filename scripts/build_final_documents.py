"""Build final academic report and editable pitch from local, reviewed evidence.
No application/provider imports, network requests, environment reads or publication.
Existing finals are copied to ignored backups before replacement. Published figures
allow reproduction without private runtime records or the original recording.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
from datetime import datetime, timezone
from pathlib import Path

import pymupdf as fitz
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'Projeto_Final_Artefatos'
SOURCE = ROOT / 'docs/InsurMinds_Relatorio_Tecnico.md'
FIGURES = OUT / 'figuras'
DIAGNOSTICS = ROOT / 'data/processed/final_release/documents'
FRAMES = ROOT / 'data/processed/final_release/video_inspection'
PDF_PATH = OUT / 'InsurMinds_Relatorio_Tecnico.pdf'
PPTX_PATH = OUT / 'InsurMinds_Projeto_Final.pptx'
ARCHITECTURE = OUT / 'InsurMinds_Arquitetura.png'
NAVY, TEAL, INK, GRAY, PALE, WHITE = '17324D', '176579', '263B4C', '607381', 'F0F5F7', 'FFFFFF'
FONT_REG = Path('C:/Windows/Fonts/arial.ttf')
FONT_BOLD = Path('C:/Windows/Fonts/arialbd.ttf')
REG = fitz.Font(fontfile=str(FONT_REG)) if FONT_REG.exists() else fitz.Font('helv')
BOLD = fitz.Font(fontfile=str(FONT_BOLD)) if FONT_BOLD.exists() else fitz.Font('hebo')
TEAM = ['José Leonardo Alves Vilela', 'Vitor Ferreira', 'Wagner Assis']
REPOSITORY = 'https://github.com/malandrindev/Apolices-Projeto-Final'


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def rgb(value: str) -> tuple[float, float, float]:
    return tuple(int(value[i:i + 2], 16) / 255 for i in (0, 2, 4))


def clean(text: str) -> str:
    return text.translate(str.maketrans({'“': '"', '”': '"', '’': "'", '‘': "'", '—': '-', '–': '-', '·': '|'}))


def wrap(text: str, width: float, size: float, bold: bool = False) -> list[str]:
    font = BOLD if bold else REG
    lines = []
    for paragraph in text.split('\n'):
        line = ''
        for word in paragraph.split():
            if font.text_length((line + ' ' + word).strip(), fontsize=size) <= width:
                line = (line + ' ' + word).strip()
                continue
            if line:
                lines.append(line)
                line = ''
            for character in word:
                if line and font.text_length(line + character, fontsize=size) > width:
                    lines.append(line)
                    line = ''
                line += character
        if line:
            lines.append(line)
    return lines


def pillow_font(size: int, bold: bool = False):
    path = FONT_BOLD if bold else FONT_REG
    if path.exists():
        return ImageFont.truetype(str(path), size)
    try:
        return ImageFont.truetype('DejaVuSans-Bold.ttf' if bold else 'DejaVuSans.ttf', size)
    except OSError:
        return ImageFont.load_default(size=size)


def architecture() -> None:
    if ARCHITECTURE.exists():
        return
    im = Image.new('RGB', (1920, 1080), '#FFFFFF')
    d = ImageDraw.Draw(im)
    d.text((78, 54), 'InsurMinds | arquitetura da solução', fill='#17324D', font=pillow_font(54, True))
    d.text((80, 134), 'Do documento à comparação verificável', fill='#607381', font=pillow_font(31))
    labels = [
        ('1', 'Documentos', 'PDF nativo, digitalizado\nou imagem'),
        ('2', 'Leitura local', 'Validação, PyMuPDF\ne Tesseract'),
        ('3', 'Organização', 'Segmentação e\nrecuperação local'),
        ('4', 'IA Generativa', 'Extração, interpretação\ne modelo por tarefa'),
        ('8', 'Interface', 'Evidências, revisão\ne exportação'),
        ('7', 'Comparação', 'Referência versus\ndocumentos candidatos'),
        ('6', 'Persistência', 'JSON, SQLite\ne cache local'),
        ('5', 'Validação', '27 campos, página,\ntrecho e qualificadores'),
    ]
    boxes = []
    for i, (n, title, body) in enumerate(labels):
        x, y = 80 + (i % 4) * 458, 254 + (i // 4) * 314
        boxes.append((x, y, x + 394, y + 226))
        d.rounded_rectangle(boxes[-1], radius=18, fill='#F0F5F7', outline='#176579', width=3)
        d.text((x + 23, y + 20), n, font=pillow_font(33, True), fill='#176579')
        d.text((x + 23, y + 73), title, font=pillow_font(33, True), fill='#17324D')
        d.multiline_text((x + 23, y + 127), body, font=pillow_font(25), fill='#263B4C', spacing=10)
    for row in (0, 1):
        for col in range(3):
            x, y = 80 + col * 458 + 394, 367 + row * 314
            if row == 0:
                d.line((x + 9, y, x + 52, y), fill='#176579', width=5)
                d.polygon([(x + 54, y), (x + 38, y - 10), (x + 38, y + 10)], fill='#176579')
            else:
                d.line((x + 9, y, x + 52, y), fill='#176579', width=5)
                d.polygon([(x + 7, y), (x + 23, y - 10), (x + 23, y + 10)], fill='#176579')
    d.line((1651, 492, 1651, 556), fill='#176579', width=5)
    d.polygon([(1651, 566), (1641, 547), (1661, 547)], fill='#176579')
    d.text((80, 907), 'Evidência primeiro. Conclusão proporcional ao que a fonte permite.', fill='#17324D', font=pillow_font(37, True))
    d.text((80, 975), 'Ausência de evidência não demonstra ausência de cobertura.', fill='#607381', font=pillow_font(30))
    im.save(ARCHITECTURE)


def ocr_card() -> None:
    im = Image.new('RGB', (1600, 960), '#FFFFFF')
    d = ImageDraw.Draw(im)
    d.text((70, 52), 'Leitura comprovada nos registros locais', fill='#17324D', font=pillow_font(48, True))
    d.text((72, 126), 'Execução real | dois recortes de quatro páginas', fill='#607381', font=pillow_font(29))
    documents = [
        ('ALLIANZ', 'IM-ALLIANZ_native.pdf', 'native', 'Texto nativo do PDF', 'f6f97cb14c3cbafec'),
        ('PORTO', 'IM-PORTO_scan_OCR.pdf', 'tesseract', 'Reconhecimento óptico local', 'ab22ff1b58e08adc5'),
    ]
    for i, (brand, name, method, caption, digest) in enumerate(documents):
        x = 70 + i * 765
        d.rounded_rectangle((x, 218, x + 695, 754), radius=22, fill='#F0F5F7', outline='#CDDCE1', width=2)
        d.text((x + 31, 252), brand, fill='#176579', font=pillow_font(30, True))
        d.text((x + 31, 305), name, fill='#17324D', font=pillow_font(29, True))
        d.text((x + 31, 371), caption, fill='#263B4C', font=pillow_font(29))
        d.text((x + 31, 429), 'extraction_method', fill='#607381', font=pillow_font(25))
        d.text((x + 31, 468), method, fill='#176579', font=pillow_font(45, True))
        for page in range(4):
            px = x + 31 + page * 153
            d.rounded_rectangle((px, 551, px + 132, 637), radius=12, fill='#FFFFFF', outline='#176579', width=2)
            d.text((px + 32, 570), str(page + 1), fill='#17324D', font=pillow_font(40, True))
        d.text((x + 31, 671), '4 páginas | identificação SHA-256', fill='#607381', font=pillow_font(24))
        d.text((x + 31, 710), digest + '…', fill='#607381', font=pillow_font(23))
    d.text((72, 811), 'A leitura local precede a extração e a comparação com IA.', fill='#17324D', font=pillow_font(31, True))
    d.text((72, 870), 'Este cenário não corresponde à análise integral dos originais de 75 e 52 páginas.', fill='#607381', font=pillow_font(26))
    im.save(FIGURES / '07_ocr_tesseract.png')


def ensure_figures(refresh: bool = False) -> None:
    OUT.mkdir(exist_ok=True)
    FIGURES.mkdir(exist_ok=True)
    specs = [
        ('report_68_clean.png', '01_upload_referencia.png', (0, 70, 1310, 800)),
        ('processing_clean.png', '02_processamento.png', (0, 70, 1310, 850)),
        ('report_298_clean.png', '03_resumo.png', (0, 115, 1310, 525)),
        ('report_300_clean.png', '04_diferencas.png', (0, 215, 1310, 760)),
        ('evidence_clean.png', '05_evidencias.png', (0, 65, 1310, 860)),
        ('ai_usage_clean.png', '06_uso_ia.png', (0, 285, 1310, 898)),
    ]
    for source, target, crop in specs:
        destination = FIGURES / target
        if refresh or not destination.exists():
            with Image.open(FRAMES / source) as im:
                im.crop(crop).save(destination)
    if refresh or not (FIGURES / '07_ocr_tesseract.png').exists():
        ocr_card()
    architecture()


def backup() -> Path:
    stamp = datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ')
    directory = ROOT / 'data/processed/final_release/backups' / ('documents_' + stamp)
    directory.mkdir(parents=True)
    entries = []
    for source in (PDF_PATH, PPTX_PATH):
        if source.exists():
            target = directory / source.name
            shutil.copy2(source, target)
            entries.append({'path': str(source.relative_to(ROOT)).replace('\\', '/'), 'sha256': sha(source), 'bytes': source.stat().st_size})
    (directory / 'manifest.json').write_text(json.dumps({'files': entries}, ensure_ascii=False, indent=2), encoding='utf-8')
    return directory


class Report:
    def __init__(self):
        self.doc = fitz.open()
        self.page = None
        self.y = 0.0
        self.toc = [[1, '1. Capa', 1]]
        self.contents = None
        self.figure_pages = []

    def new(self, landscape=False):
        self.page = self.doc.new_page(width=841.9 if landscape else 595.3, height=595.3 if landscape else 841.9)
        for name, font in (('Body', REG), ('Strong', BOLD)):
            self.page.insert_font(fontname=name, fontbuffer=font.buffer, set_simple=True)
        w = self.page.rect.width
        self.page.draw_rect(fitz.Rect(0, 0, w, 8), fill=rgb(TEAL), color=rgb(TEAL))
        self.at(45, 31, 'INSURMINDS  |  PROJETO FINAL I2A2', 8.5, True, NAVY)
        self.at(45, 47, 'Plataforma Inteligente para Análise e Comparação de Apólices D&O', 8, color=GRAY)
        self.y = 71

    def at(self, x, y, text, size=10.7, bold=False, color=INK):
        self.page.insert_text((x, y), clean(text), fontname='Strong' if bold else 'Body', fontsize=size, color=rgb(color))

    def ensure(self, height):
        if self.page is None or self.page.rect.width > 600 or self.y + height > 790:
            self.new()

    def cover(self):
        self.new()
        self.page.draw_rect(fitz.Rect(45, 159, 550, 164), fill=rgb(TEAL), color=rgb(TEAL))
        self.at(45, 132, 'InsurMinds', 46, True, NAVY)
        y = 218
        for line in wrap('Plataforma Inteligente para Análise e Comparação de Apólices D&O', 497, 29, True):
            self.at(45, y, line, 29, True, NAVY)
            y += 38
        self.at(45, y + 36, 'Relatório Técnico', 23, True, TEAL)
        self.at(45, y + 67, 'Projeto Final I2A2 / InsurMinds', 16, color=GRAY)
        self.at(45, 538, 'EQUIPE', 10, True, TEAL)
        for i, name in enumerate(TEAM):
            self.at(45, 570 + i * 29, name, 15, color=NAVY)
        self.at(45, 704, 'Outubro de 2026', 13, True, NAVY)
        self.at(45, 735, 'Código-fonte e materiais da entrega:', 10, color=GRAY)
        self.at(45, 756, REPOSITORY, 10, color=TEAL)
        self.page.insert_link({'kind': fitz.LINK_URI, 'from': fitz.Rect(45, 744, 550, 762), 'uri': REPOSITORY})
        self.new()
        self.contents = self.page.number
        self.at(45, 100, 'Sumário', 26, True, NAVY)
        self.new()

    def heading(self, title):
        lines = wrap(title, 505, 19, True)
        self.ensure(270 + len(lines) * 24)
        self.toc.append([1, title, len(self.doc)])
        for line in lines:
            self.at(45, self.y + 19, line, 19, True, NAVY)
            self.y += 25
        self.y += 14

    def paragraph(self, text, reference=False):
        size = 9.8 if reference else 10.7
        lines = wrap(text, 505, size)
        leading = 1.42 if reference else 1.47
        self.ensure(len(lines) * size * leading + 15)
        top = self.y
        for line in lines:
            self.at(45, self.y + size, line, size)
            self.y += size * leading
        for match in re.finditer(r'https?://[^\s]+', text):
            uri = match.group(0).rstrip('.,;')
            self.page.insert_link({'kind': fitz.LINK_URI, 'from': fitz.Rect(45, top, 550, self.y), 'uri': uri})
        self.y += 7 if reference else 12

    def table(self, headers, rows):
        widths = [194, 311]
        size = 9.6
        def render(row, head=False, shaded=False):
            lines = [wrap(cell, w - 18, size, head) for cell, w in zip(row, widths)]
            height = max(map(len, lines)) * 14 + 19
            if self.page is None or self.page.rect.width > 600 or self.y + height > 790:
                self.new()
                if not head:
                    render(headers, True)
            x = 45
            for i, w in enumerate(widths):
                self.page.draw_rect(fitz.Rect(x, self.y, x + w, self.y + height), fill=rgb(NAVY if head else PALE if shaded else WHITE), color=rgb('D8E3E7'), width=.5)
                for j, line in enumerate(lines[i]):
                    self.at(x + 9, self.y + 15 + j * 14, line, size, head, WHITE if head else INK)
                x += w
            self.y += height
        self.ensure(100)
        render(headers, True)
        for i, row in enumerate(rows):
            render(row, shaded=i % 2 == 0)
        self.y += 17

    def figure(self, caption, path):
        self.new(landscape=True)
        lines = wrap(caption, 751, 12, True)
        for line in lines:
            self.at(45, self.y + 12, line, 12, True, NAVY)
            self.y += 18
        self.y += 14
        self.page.insert_image(fitz.Rect(45, self.y, 797, 548), filename=str(path), keep_proportion=True)
        self.figure_pages.append(self.page.number + 1)
        self.y = 795

    def finish(self):
        current = self.page
        self.page = self.doc[self.contents]
        y = 142
        for _, title, number in self.toc:
            lines = wrap(title, 450, 10.2)
            for line in lines:
                self.at(45, y, line, 10.2, color=INK)
                y += 16
            self.at(525, y - 16, str(number), 10.2, color=TEAL)
            self.page.insert_link({'kind': fitz.LINK_GOTO, 'from': fitz.Rect(45, y - len(lines) * 16 - 11, 550, y + 3), 'page': number - 1})
            y += 6
        self.page = current
        for index, page in enumerate(self.doc):
            w, h = page.rect.width, page.rect.height
            page.draw_line((45, h - 31), (w - 45, h - 31), color=rgb('CCD9DF'), width=.6)
            page.insert_text((45, h - 16), 'Relatório Técnico | MVP acadêmico | Outubro de 2026', fontname='Body', fontsize=7.8, color=rgb(GRAY))
            page.insert_text((w - 92, h - 16), f'{index + 1} / {len(self.doc)}', fontname='Body', fontsize=7.8, color=rgb(GRAY))
        self.doc.set_toc(self.toc)
        self.doc.set_metadata({'title': 'InsurMinds — Relatório Técnico', 'author': '; '.join(TEAM), 'subject': 'MVP acadêmico de análise e comparação documental D&O', 'creator': 'scripts/build_final_documents.py'})
        self.doc.save(PDF_PATH, garbage=4, deflate=True)
        return {'pages': len(self.doc), 'figure_pages': self.figure_pages, 'sections': len(self.toc)}


def report():
    doc = Report()
    doc.cover()
    lines = SOURCE.read_text(encoding='utf-8').splitlines()
    i, active, reference = 0, False, False
    while i < len(lines):
        line = lines[i].strip()
        if line.startswith('## '):
            title = line[3:]
            active = not title.startswith('1. ')
            reference = title.startswith('24. ')
            if active:
                doc.heading(title)
            i += 1
            continue
        if not active or not line:
            i += 1
            continue
        image = re.fullmatch(r'!\[(.+)]\((.+)\)', line)
        if image:
            doc.figure(image.group(1), (SOURCE.parent / image.group(2)).resolve())
            i += 1
            continue
        if line.startswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append([cell.strip() for cell in lines[i].strip().strip('|').split('|')])
                i += 1
            doc.table(table_lines[0], table_lines[2:])
            continue
        paragraph = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith(('## ', '![', '|')):
            paragraph.append(lines[i].strip())
            i += 1
        doc.paragraph(' '.join(paragraph), reference)
    return doc.finish()


def color(value):
    return RGBColor.from_string(value)


def text(slide, x, y, w, h, value, size=22, bold=False, ink=INK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.word_wrap = True
    for i, line in enumerate(value.split('\n')):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.font.name = 'Arial'
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color(ink)
        p.space_after = Pt(9)
    return box


def rect(slide, x, y, w, h, fill=PALE, border=None, radius=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    if border:
        shape.line.color.rgb = color(border)
    else:
        shape.line.fill.background()
    return shape


def picture(slide, path, x, y, w, h):
    with Image.open(path) as im:
        ratio = im.width / im.height
    actual_w = min(w, h * ratio)
    actual_h = actual_w / ratio
    slide.shapes.add_picture(str(path), Inches(x + (w - actual_w) / 2), Inches(y + (h - actual_h) / 2), width=Inches(actual_w), height=Inches(actual_h))


def pitch():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    def new(title, subtitle=None):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        rect(s, 0, 0, 13.333, .09, TEAL)
        text(s, .62, .39, 12.1, .63, title, 30, True, NAVY)
        if subtitle:
            text(s, .64, 1.09, 12, .48, subtitle, 17, ink=GRAY)
        text(s, .64, 7.15, 11.6, .22, 'InsurMinds | Projeto Final I2A2 | Outubro de 2026', 10, ink=GRAY)
        text(s, 12.13, 7.13, .6, .26, f'{len(prs.slides):02d}', 12, ink=TEAL)
        return s

    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, 13.333, 7.5, NAVY)
    rect(s, .7, 1.47, 1.0, .09, TEAL)
    text(s, .7, .67, 12, .9, 'InsurMinds', 48, True, WHITE)
    text(s, .72, 2.0, 11.65, 1.85, 'Análise e comparação de\ndocumentos D&O com evidências', 34, True, WHITE)
    text(s, .75, 4.25, 11.5, .65, 'Projeto Final I2A2 / InsurMinds', 22, ink='B9DCE2')
    text(s, .75, 5.25, 11.75, 1.2, 'José Leonardo Alves Vilela\nVitor Ferreira  |  Wagner Assis', 19, ink=WHITE)
    text(s, .75, 6.74, 11.75, .28, 'Outubro de 2026', 15, ink='B9DCE2')

    s = new('O problema: comparar exige preservar o contexto', 'Coberturas e limites dependem de definições, condições, sujeitos e exceções.')
    items = [('Localizar', 'Documentos extensos dispersam informações relevantes.'), ('Interpretar', 'Termos semelhantes podem produzir alcances diferentes.'), ('Conferir', 'Toda conclusão precisa voltar à página e ao trecho de origem.')]
    for i, (title, body) in enumerate(items):
        x = .66 + 4.23 * i
        rect(s, x, 2.01, 3.86, 3.31, PALE, radius=True)
        text(s, x + .24, 2.3, 3.38, .5, title, 27, True, TEAL)
        text(s, x + .24, 3.07, 3.3, 1.87, body, 23)
    text(s, .7, 5.88, 11.9, .71, 'A plataforma organiza a análise e conserva a evidência para o especialista.', 23, True, NAVY)

    s = new('Uma jornada simples para uma análise verificável', 'Adicionar documentos → escolher a referência → comparar → conferir evidências.')
    picture(s, FIGURES / '01_upload_referencia.png', .65, 1.77, 8.67, 4.8)
    rect(s, 9.59, 1.97, 3.09, 4.39, PALE, radius=True)
    text(s, 9.82, 2.26, 2.65, 3.55, '2 a 5 documentos\nReferência explícita\nComparação por pares\nExportação do resultado', 21, True, NAVY)

    s = new('Arquitetura com responsabilidades delimitadas')
    picture(s, ARCHITECTURE, .48, 1.25, 12.37, 5.68)

    s = new('PDF nativo e digitalizado no mesmo fluxo', 'A leitura local acontece antes da extração e da comparação com IA.')
    picture(s, FIGURES / '07_ocr_tesseract.png', .56, 1.76, 7.53, 4.75)
    picture(s, FIGURES / '02_processamento.png', 8.25, 1.78, 4.52, 3.55)
    text(s, 8.48, 5.51, 4.07, .83, '4 páginas nativas\n4 páginas com Tesseract', 20, True, TEAL)

    s = new('Execução real: resultado proporcional à evidência', 'Allianz e Porto | recortes de 4 + 4 páginas | fluxo completo registrado.')
    picture(s, FIGURES / '03_resumo.png', .62, 1.91, 12.05, 3.89)
    text(s, .77, 5.99, 11.8, .68, '2 diferenças  |  0 equivalências  |  19 itens sem conclusão segura', 23, True, NAVY)

    s = new('Da diferença ao documento que a sustenta', 'Documento, página e trecho podem ser conferidos na própria interface.')
    picture(s, FIGURES / '05_evidencias.png', .6, 1.77, 8.8, 4.87)
    rect(s, 9.65, 1.91, 3.0, 4.55, PALE, radius=True)
    text(s, 9.86, 2.22, 2.58, 3.78, 'Citação verificável\n\nContexto preservado\n\nRevisão humana\n\nSem escolha global automática', 20, True, NAVY)

    s = new('Transparência sobre o uso de IA', 'Acompanhamento opcional de modelos, chamadas e tokens.')
    picture(s, FIGURES / '06_uso_ia.png', .58, 1.81, 8.7, 4.79)
    text(s, 9.6, 2.09, 3.02, 1.8, '37 chamadas na execução histórica\n\n19 + 17 extração\n1 comparação', 21, True, NAVY)
    text(s, 9.6, 4.35, 3.02, 1.72, 'Cache do provedor e cache local são distintos.\nTokens não comprovam cobrança efetiva.', 18, ink=GRAY)

    s = new('Limites assumidos e revisão responsável', 'Conclusões conservadoras fazem parte do resultado.')
    items = [
        ('Fonte', 'Condições gerais não comprovam a contratação individual.'),
        ('Evidência', 'Recortes não representam análise integral de 75 e 52 páginas.'),
        ('Interpretação', 'OCR e modelos podem omitir detalhes; o especialista deve conferir.'),
        ('Demonstração', 'Especificações fictícias são identificadas e não têm validade contratual.'),
    ]
    for i, (title, body) in enumerate(items):
        x, y = .68 + (i % 2) * 6.38, 1.96 + (i // 2) * 2.31
        rect(s, x, y, 5.99, 2.0, PALE, radius=True)
        text(s, x + .22, y + .19, 5.5, .42, title, 23, True, TEAL)
        text(s, x + .22, y + .79, 5.5, 1.06, body, 20)

    s = new('Uma base funcional para evolução', 'Automação útil exige rastreabilidade e prudência.')
    text(s, .72, 1.95, 11.8, 1.25, 'O InsurMinds integra OCR, IA, estruturação e comparação\nem uma jornada de apoio à análise documental.', 27, True, NAVY)
    text(s, .76, 3.62, 11.8, 1.82, 'Próximas evoluções\nAvaliação com especialistas e documentos de escopo semelhante\nMelhor tratamento de continuidade, tabelas, endossos e conflitos', 22)
    text(s, .76, 6.02, 11.8, .45, REPOSITORY, 18, ink=TEAL)
    prs.core_properties.title = 'InsurMinds — Projeto Final I2A2'
    prs.core_properties.subject = 'Plataforma de análise e comparação de documentos D&O'
    prs.core_properties.author = '; '.join(TEAM)
    prs.core_properties.comments = 'Capturas reais da interface final. Cenário real de 4 + 4 páginas; demonstração fictícia identificada. Sem novas APIs na geração.'
    prs.save(PPTX_PATH)
    return {'slides': len(prs.slides), 'editable_text': True}


def render_pdf(directory: Path):
    doc = fitz.open(PDF_PATH)
    directory.mkdir(parents=True, exist_ok=True)
    rows = []
    for i, page in enumerate(doc):
        path = directory / f'page_{i + 1:02}.png'
        page.get_pixmap(matrix=fitz.Matrix(1.4, 1.4), alpha=False).save(path)
        escaped = []
        for block in page.get_text('dict')['blocks']:
            if block['type'] != 0:
                continue
            for line in block['lines']:
                for span in line['spans']:
                    b = fitz.Rect(span['bbox'])
                    if b.x0 < 0 or b.y0 < 0 or b.x1 > page.rect.width + 1 or b.y1 > page.rect.height + 1:
                        escaped.append(span['text'])
        rows.append({'page': i + 1, 'width': page.rect.width, 'height': page.rect.height, 'outside_page_text': escaped, 'text_chars': len(page.get_text()), 'image': str(path.relative_to(ROOT)).replace('\\', '/')})
    thumbs = []
    for row in rows:
        with Image.open(ROOT / row['image']) as im:
            canvas = Image.new('RGB', (310, 455), '#E4EBEE')
            im.thumbnail((300, 413))
            canvas.paste(im, ((310 - im.width) // 2, 20))
            ImageDraw.Draw(canvas).text((10, 432), f"Página {row['page']}", fill='#17324D', font=pillow_font(16, True))
            thumbs.append(canvas)
    for start in range(0, len(thumbs), 8):
        subset = thumbs[start:start + 8]
        sheet = Image.new('RGB', (1240, 910), '#E4EBEE')
        for i, im in enumerate(subset):
            sheet.paste(im, ((i % 4) * 310, (i // 4) * 455))
        sheet.save(directory / f'contact_{start + 1:02}_{start + len(subset):02}.jpg', quality=92)
    return rows


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--refresh-images', action='store_true', help='Recreate reviewed figures from local original recording frames.')
    args = parser.parse_args()
    ensure_figures(args.refresh_images)
    saved = backup()
    pdf = report()
    deck = pitch()
    DIAGNOSTICS.mkdir(parents=True, exist_ok=True)
    qa = render_pdf(DIAGNOSTICS / 'report_pages')
    result = {
        'new_provider_calls': 0,
        'source': str(SOURCE.relative_to(ROOT)).replace('\\', '/'),
        'source_sha256': sha(SOURCE),
        'report': {**pdf, 'sha256': sha(PDF_PATH), 'bytes': PDF_PATH.stat().st_size},
        'pitch': {**deck, 'sha256': sha(PPTX_PATH), 'bytes': PPTX_PATH.stat().st_size},
        'backup_directory': str(saved.relative_to(ROOT)).replace('\\', '/'),
        'figures': [{'path': str(p.relative_to(ROOT)).replace('\\', '/'), 'sha256': sha(p)} for p in [ARCHITECTURE, *sorted(FIGURES.glob('*.png'))]],
        'pdf_layout': qa,
        'visual_review_required': True,
    }
    (DIAGNOSTICS / 'build_manifest.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({k: result[k] for k in ('new_provider_calls', 'report', 'pitch', 'backup_directory')}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
